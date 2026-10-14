from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


wb = Workbook()
ws = wb.active

ws['A1'] = "Prodotto"
ws['A1'].font = Font(bold=True, color="FF0090")
ws['A1'].fill = PatternFill("solid", fgColor="CC5200")

ws['A1'].alignment = Alignment(horizontal="center")

wb.save("./excelpy/stile.xlsx")