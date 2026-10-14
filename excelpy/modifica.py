from openpyxl import load_workbook

wb = load_workbook("./excelpy/file_2.xlsx")
ws = wb.active

ws["A2"] = "Anguria"
ws["B2"] = 5.99

ws.append(["Mandarino", 5.54, 125])

wb.save("./excelpy/file_2.xlsx")
