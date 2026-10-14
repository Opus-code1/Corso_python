from openpyxl import load_workbook

wb = load_workbook("./excelpy/file_2.xlsx")

ws = wb.active

dati = ws["A1"].value
print(dati)

# avanzato 
for riga in ws.iter_rows(values_only=True):
    print(riga)


