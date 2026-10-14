from openpyxl import Workbook

wb = Workbook()
ws1 = wb.active

ws1.title = "Lunedì"

ws2 = wb.create_sheet("Martedì")
ws3 = wb.create_sheet("Mercoledì")

ws1['A1'] = "Spesa"
ws2['A1'] = "Concerto"
ws3['A1'] = "Cena"


wb.save("./excelpy/multifoglio.xlsx")