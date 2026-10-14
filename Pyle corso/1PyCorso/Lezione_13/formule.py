from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

wb = Workbook()
ws = wb.active

ws.title = "Formule"

# ws["A1"] = 10
# ws["A2"] = 20
# ws["A3"] = "=SUM(A1:A2)"

ws["A1"] = "Nome"
ws["B1"] = "Età"
ws["C1"] = "Città"

data = [
    ["Marco", 25, "Reggio Emilia"],
    ["Lisa", 28, "Reggio Calabria"],
    ["Sarah", 21, "Palermo"],
    ["Giulia", 24, "Catania"]
]

for riga in data:
    ws.append(riga)

row_num = len(data) + 3

tab = Table(displayName="Nome", ref=f"B1:C{row_num}")
style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True,showColumnStripes=True,showFirstColumn=False,showLastColumn=False)

tab.tableStyleInfo = style
ws.add_table(tab)

wb.save("esempio_1.xlsx")
print("salvato")