import matplotlib.pyplot as plt
from openpyxl import load_workbook

file = './MatPlot/vendite.xlsx'

prodotti, prezzi, quantita, totali = [],[],[],[]

wb = load_workbook(file)
ws = wb["Vendite"]

for row in ws.iter_rows(min_row=2, values_only=True):
    if row[0] is None:
        continue
    prodotti.append(row[0])
    prezzi.append(row[1])
    quantita.append(row[2])
    totali.append(row[3])

print(prodotti,prezzi,quantita,totali)

top10_idx = sorted(range(len(totali)), key=lambda i: totali[i], reverse=True)[:10]
top10_prod = [prodotti[i] for i in top10_idx]
top10_val = [totali[i] for i in top10_idx]

plt.figure(figsize=(12, 6))
plt.title("=== Top 10 prodotti ===")
plt.barh(top10_prod, top10_val, color="blue")
plt.xlabel("Testo per la X")
plt.ylabel("Testo per la Y")
plt.tight_layout()
plt.show()