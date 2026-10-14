# === EXCEL BOT === #
# Leggere file Excel
# Calcolare i  totali
# Aggiungere formattazione
# Salvare il file

# =====================================================
# IMPORTAZIONI
# =====================================================

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os

# === FUNZIONI === #
'la funzione che permette di leggere i dati'

def leggi_vendite(percorso):
    if not os.path.exists(percorso):
        print(f"ERRORE: il file {percorso} non esiste")

    wb = load_workbook(percorso)

    ws = wb.active
    
    dati = []
    
    for riga in ws.iter_rows(min_row=2, values_only=True):
        if riga[0] is not None:
            dati.append(riga)
    return dati

def calcolo_totali(dati):
    risultati = []
    
    for riga in dati:
        prodotto = riga[0]
        
        prezzo = riga[1]
        
        quantita = riga[2]
        
        totale = prezzo * quantita
        
        nuova_riga = (prodotto, prezzo, quantita, totale)
        
        risultati.append(nuova_riga)
    return risultati

def crea_report(dati_con_tot, percorso_output):
    wb = load_workbook("./ExcelBot/vendite.xlsx")
    ws = wb.active
    
    ws.title = "Report"
    
    intestazione = ['Prodotto','Prezzo','Quantità','Totale']
    ws.append(intestazione)
    
    colore_sfondo = "3F41D4"
    colore_txt = "FFFFFF"
    riempimento = PatternFill("solid", fgColor=colore_sfondo)
    
    font_intest = Font(bold=True, color=colore_txt)
    allineamento = Alignment(horizontal="center")
    
    for cella in ws[1]:
        cella.font = font_intest
        cella.fill = riempimento
        cella.alignment = allineamento
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 12
        
        for riga in dati_con_tot:
            ws.append(riga)
            
        ultima_riga = ws.max_row
        
        riga_totale = ultima_riga + 1
        
        ws[f"A{riga_totale}"] = "TOTALE"
        
        ws[f"D{riga_totale}"] = f"=SUM(D2:D{ultima_riga})"
        
        ws[f"A{riga_totale}"].font = Font(bold=True)
        
        ws[f"D{riga_totale}"].font = Font(bold=True)
        
        wb.save(percorso_output)
        print(f"File salvato in: {percorso_output}")


file_input = "./ExcelBot/vendite.xlsx"
file_output = "./ExcelBot/report_vendite.xlsx"

dati = leggi_vendite(file_input)

if dati is None:
    print("Nessun dato è stato trovato")
else:
    print(f"Lette {len(dati)} righe di file")
    
    dati_con_totali = calcolo_totali(dati)
    
    crea_report(dati_con_totali, file_output)
    